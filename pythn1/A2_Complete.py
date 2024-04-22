from docx import Document
from openpyxl import Workbook
import os

# Get list of documents
dirName = "C:\\Users\\balvi\\mycode\\sem2\\python\\Assigntwo\\"
dirList = os.listdir(dirName)
docList = []
for f in dirList:
    if os.path.isfile(os.path.join(dirName, f)):
        if f.split('.')[1] == "docx":
            docList.append(f)

# Create invoice data
invoices = []
for doc in docList:
    aDoc = Document(os.path.join(dirName, doc))
    invoiceNum = doc.split('.')[0]
    totalQty = 0

    for p in aDoc.paragraphs:
        body = p.text

        # Tally product quantities
        if "PRODUCTS" in body:
            products = body.split("\n")
            for i in range(1, len(products)-1):
                totalQty += int(products[i].split(':')[1])

        # Get price info
        elif "SUBTOTAL" in body:
            priceLines = body.split("\n")
            subtot = priceLines[0].split(':')[1]
            tax = priceLines[1].split(':')[1]
            total = priceLines[2].split(':')[1]

    # Put the retrieved info into a dictionary and add to the invoice list
    invoices.append({"Invoice Number": invoiceNum, "Total Quantity": int(totalQty), "Subtotal": float(subtot), "Tax": float(tax), "Total": float(total)})

# Create a workbook and populate with the invoice data
wb = Workbook()
st = wb.active
headers = ["Invoice Number", "Total Quantity", "Subtotal", "Tax", "Total"]
cols = ['A', 'B', 'C', 'D', 'E']
rows = range(0, len(invoices))

# Add headers
for i in range(len(cols)):
    st[f"{cols[i]}1"] = headers[i]

for i in rows:
    invoice = invoices[i]
    for j in range(len(cols)):
        # Rows offset by 2 (1 row b/c sheets are not 0-indexed, 1 row because of our added header row)
        st[f"{cols[j]}{i+2}"] = invoice[headers[j]]

wb.save("A2_Complete.xlsx")
